"""
Модуль генерации Word-документов.
Создаёт .docx файлы из сгенерированного текста с правильным оформлением.
"""

import logging
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

from utils.text_utils import parse_generated_text, sanitize_filename

logger = logging.getLogger(__name__)


def _setup_styles(doc: Document, template_id: str) -> None:
    """Настройка стилей документа в зависимости от шаблона."""

    # Стиль обычного текста (Normal)
    style_normal = doc.styles["Normal"]
    font = style_normal.font
    font.name = "Times New Roman"
    font.size = Pt(14)
    font.color.rgb = RGBColor(0, 0, 0)
    paragraph_format = style_normal.paragraph_format
    paragraph_format.space_after = Pt(0)
    paragraph_format.space_before = Pt(0)
    paragraph_format.line_spacing = 1.5
    paragraph_format.first_line_indent = Cm(1.25)
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Стиль Heading 1
    if "Heading 1" in doc.styles:
        h1 = doc.styles["Heading 1"]
    else:
        h1 = doc.styles.add_style("Heading 1", WD_STYLE_TYPE.PARAGRAPH)
    h1_font = h1.font
    h1_font.name = "Times New Roman"
    h1_font.size = Pt(16)
    h1_font.bold = True
    h1_font.color.rgb = RGBColor(0, 0, 0)
    h1_pf = h1.paragraph_format
    h1_pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1_pf.space_before = Pt(12)
    h1_pf.space_after = Pt(12)
    h1_pf.line_spacing = 1.5
    h1_pf.first_line_indent = Cm(0)

    # Стиль Heading 2
    if "Heading 2" in doc.styles:
        h2 = doc.styles["Heading 2"]
    else:
        h2 = doc.styles.add_style("Heading 2", WD_STYLE_TYPE.PARAGRAPH)
    h2_font = h2.font
    h2_font.name = "Times New Roman"
    h2_font.size = Pt(15)
    h2_font.bold = True
    h2_font.color.rgb = RGBColor(0, 0, 0)
    h2_pf = h2.paragraph_format
    h2_pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2_pf.space_before = Pt(10)
    h2_pf.space_after = Pt(6)
    h2_pf.line_spacing = 1.5
    h2_pf.first_line_indent = Cm(0)

    # Стиль Heading 3
    if "Heading 3" in doc.styles:
        h3 = doc.styles["Heading 3"]
    else:
        h3 = doc.styles.add_style("Heading 3", WD_STYLE_TYPE.PARAGRAPH)
    h3_font = h3.font
    h3_font.name = "Times New Roman"
    h3_font.size = Pt(14)
    h3_font.bold = True
    h3_font.color.rgb = RGBColor(0, 0, 0)
    h3_pf = h3.paragraph_format
    h3_pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3_pf.space_before = Pt(8)
    h3_pf.space_after = Pt(4)
    h3_pf.line_spacing = 1.5
    h3_pf.first_line_indent = Cm(0)

    # Настройка полей страницы
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    if template_id == "university":
        # Для вузовского реферата — чуть другие поля
        for section in doc.sections:
            section.left_margin = Cm(3)
            section.right_margin = Cm(1)


def create_document(text: str, template_id: str, topic: str, output_dir: Path) -> Path:
    """
    Создаёт Word-документ из сгенерированного текста.

    Args:
        text: Сгенерированный текст реферата (markdown)
        template_id: ID шаблона
        topic: Тема реферата
        output_dir: Директория для сохранения

    Returns:
        Путь к созданному .docx файлу
    """
    # Парсим markdown в структуру
    elements = parse_generated_text(text)

    if not elements:
        raise ValueError("Не удалось распарсить сгенерированный текст")

    # Создаём документ
    doc = Document()
    _setup_styles(doc, template_id)

    # Заполняем документ
    for element in elements:
        el_type = element["type"]
        el_text = element["text"]

        if el_type == "heading1":
            para = doc.add_paragraph(el_text)
            para.style = doc.styles["Heading 1"]

        elif el_type == "heading2":
            para = doc.add_paragraph(el_text)
            para.style = doc.styles["Heading 2"]

        elif el_type == "heading3":
            para = doc.add_paragraph(el_text)
            para.style = doc.styles["Heading 3"]

        elif el_type == "list_item":
            para = doc.add_paragraph(el_text, style="List Bullet")
            # Настраиваем шрифт элемента списка
            for run in para.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(14)

        elif el_type == "paragraph":
            para = doc.add_paragraph(el_text)
            para.style = doc.styles["Normal"]

    # Формируем имя файла
    safe_name = sanitize_filename(topic)
    filename = f"Реферат_{safe_name}.docx"
    output_path = output_dir / filename

    # Создаём директорию если не существует
    output_dir.mkdir(parents=True, exist_ok=True)

    # Сохраняем
    doc.save(str(output_path))
    logger.info(f"Документ сохранён: {output_path}")

    return output_path
